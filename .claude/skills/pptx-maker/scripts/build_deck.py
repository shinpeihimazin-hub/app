#!/usr/bin/env python3
"""deck.json をテンプレ(.pptx)のレイアウトに流し込んで .pptx を生成する（S5）。

usage: build_deck.py template.pptx deck.json out.pptx
deck.json の形式は references/draft-format.md 参照。
"""
import argparse
import copy
import json
import sys

from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import PP_PLACEHOLDER


def find_layout(prs, key):
    layouts = list(prs.slide_layouts)
    if isinstance(key, int):
        if 0 <= key < len(layouts):
            return layouts[key]
        raise KeyError(f"レイアウト番号 {key} は存在しない（0〜{len(layouts)-1}）")
    for lay in layouts:
        if lay.name == key:
            return lay
    names = [l.name for l in layouts]
    raise KeyError(f"レイアウト名 '{key}' は存在しない。実在: {names}")


def set_body(slide, items):
    """BODY/CONTENT系の最初のプレースホルダに箇条書きを流し込む。ネスト=レベル+1。"""
    body_types = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY}
    target = None
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == 0:  # タイトルは除外
            continue
        if ph.placeholder_format.type in body_types or ph.has_text_frame:
            target = ph
            break
    if target is None:
        raise KeyError("本文プレースホルダが見つからない（layoutの選択を見直す）")
    tf = target.text_frame
    tf.clear()
    first = True

    def add(text, level):
        nonlocal first
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.text = str(text)
        p.level = level

    for item in items:
        if isinstance(item, list):
            for sub in item:
                add(sub, 1)
        else:
            add(item, 0)


def add_images(slide, images):
    for im in images:
        if "placeholder" in im:
            ph = next((p for p in slide.placeholders
                       if p.placeholder_format.idx == im["placeholder"]), None)
            if ph is None:
                raise KeyError(f"画像プレースホルダ idx={im['placeholder']} が無い")
            ph.insert_picture(im["path"])
        else:
            kw = {"left": Inches(im["left_in"]), "top": Inches(im["top_in"])}
            if "width_in" in im:
                kw["width"] = Inches(im["width_in"])
            if "height_in" in im:
                kw["height"] = Inches(im["height_in"])
            slide.shapes.add_picture(im["path"], **kw)


def build(template, deck, out):
    prs = Presentation(template)
    # テンプレに残っている既存スライドは触らない（末尾に追加していく）
    for i, sd in enumerate(deck["slides"], 1):
        try:
            layout = find_layout(prs, sd["layout"])
            slide = prs.slides.add_slide(layout)
            if sd.get("title") is not None and slide.shapes.title is not None:
                slide.shapes.title.text = sd["title"]
            if sd.get("body"):
                set_body(slide, sd["body"])
            for idx, text in (sd.get("placeholders") or {}).items():
                ph = next((p for p in slide.placeholders
                           if p.placeholder_format.idx == int(idx)), None)
                if ph is None:
                    raise KeyError(f"プレースホルダ idx={idx} が無い")
                ph.text_frame.text = str(text)
            if sd.get("images"):
                add_images(slide, sd["images"])
            if sd.get("notes"):
                slide.notes_slide.notes_text_frame.text = sd["notes"]
        except Exception as e:
            raise RuntimeError(f"スライド{i}（{sd.get('title', '無題')}）: {e}") from e
    if deck.get("title"):
        prs.core_properties.title = deck["title"]
    prs.save(out)
    return len(deck["slides"])


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("deck_json")
    ap.add_argument("out")
    args = ap.parse_args()
    with open(args.deck_json, encoding="utf-8") as fp:
        deck = json.load(fp)
    n = build(args.template, deck, args.out)
    print(f"生成完了: {args.out}（{n}スライド）")


if __name__ == "__main__":
    sys.exit(main())
