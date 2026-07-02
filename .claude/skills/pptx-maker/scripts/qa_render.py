#!/usr/bin/env python3
"""生成済み .pptx の品質QA（S6）: 幾何チェック＋レンダリング。

usage: qa_render.py deck.pptx outdir/
出力:
  outdir/qa_report.json  … 機械警告（文字あふれ推定・空プレースホルダ）
  outdir/<name>.pdf      … LibreOffice(soffice)がある場合のみ。無ければスキップと明記
  outdir/slide-XX.png    … pypdfium2がある場合のみ（PDF経由）
"""
import argparse
import json
import os
import shutil
import subprocess
import sys

from pptx import Presentation
from pptx.util import Emu

# 経験的な密度上限（全角相当・文字/平方インチ）。超えたら「あふれ疑い」警告
DENSITY_LIMIT = 18.0


def geometry_checks(path):
    warns = []
    prs = Presentation(path)
    slides = list(prs.slides)
    for si, slide in enumerate(slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text.strip()
            is_ph = sh.is_placeholder
            if is_ph and not text:
                warns.append({"slide": si, "kind": "empty_placeholder",
                              "detail": f"空のプレースホルダ: {sh.name}"})
                continue
            if not text or sh.width is None or sh.height is None:
                continue
            area = Emu(sh.width).inches * Emu(sh.height).inches
            if area <= 0:
                continue
            # 全角=1、半角=0.5 で文字量を概算
            units = sum(1.0 if ord(c) > 0xFF else 0.5 for c in text)
            density = units / area
            if density > DENSITY_LIMIT:
                warns.append({"slide": si, "kind": "overflow_risk",
                              "detail": f"{sh.name}: 密度{density:.1f}字/in²"
                                        f"（上限{DENSITY_LIMIT}）文字あふれ疑い"})
    return warns, len(prs.slides.__iter__.__self__._sldIdLst)


def find_soffice():
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    return None


def render_pdf(soffice, pptx, outdir):
    env = os.environ.copy()
    cmd = [soffice, "-env:UserInstallation=file:///tmp/lo_profile_qa",
           "--headless", "--convert-to", "pdf", "--outdir", outdir, pptx]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=180, env=env)
    pdf = os.path.join(outdir, os.path.splitext(os.path.basename(pptx))[0] + ".pdf")
    return pdf if os.path.exists(pdf) else None


def render_pngs(pdf, outdir):
    try:
        import pypdfium2 as pdfium
    except ImportError:
        return 0
    doc = pdfium.PdfDocument(pdf)
    for i, page in enumerate(doc, 1):
        img = page.render(scale=1.5).to_pil()
        img.save(os.path.join(outdir, f"slide-{i:02d}.png"))
    return len(doc)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("outdir")
    args = ap.parse_args()
    os.makedirs(args.outdir, exist_ok=True)

    warns, _ = geometry_checks(args.pptx)
    report = {"pptx": args.pptx, "warnings": warns,
              "render": {"pdf": None, "png_pages": 0, "note": ""}}

    soffice = find_soffice()
    if soffice:
        pdf = render_pdf(soffice, args.pptx, args.outdir)
        if pdf:
            report["render"]["pdf"] = pdf
            report["render"]["png_pages"] = render_pngs(pdf, args.outdir)
        else:
            report["render"]["note"] = "sofficeはあるがPDF変換失敗（Impressモジュール欠落の可能性）"
    else:
        report["render"]["note"] = "soffice未導入のためレンダリング省略（幾何チェックのみ）"

    rp = os.path.join(args.outdir, "qa_report.json")
    with open(rp, "w", encoding="utf-8") as fp:
        json.dump(report, fp, ensure_ascii=False, indent=2)

    print(f"警告 {len(warns)} 件 / PDF: {report['render']['pdf'] or '無'} "
          f"/ PNG {report['render']['png_pages']}枚 → {rp}")
    for w in warns:
        print(f"  [slide {w['slide']}] {w['kind']}: {w['detail']}")


if __name__ == "__main__":
    sys.exit(main())
