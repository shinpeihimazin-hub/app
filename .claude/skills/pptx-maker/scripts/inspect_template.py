#!/usr/bin/env python3
"""テンプレ(.pptx)のスライドマスター/レイアウト/プレースホルダを列挙する（T1: テンプレ地図）。

usage: inspect_template.py template.pptx [--json template-map.json]
"""
import argparse
import json
import sys

try:  # `| head` 等でパイプが閉じても静かに終了する（Windowsには SIGPIPE が無い）
    from signal import signal, SIGPIPE, SIG_DFL
    signal(SIGPIPE, SIG_DFL)
except ImportError:
    pass

from pptx import Presentation
from pptx.util import Emu


def emu_in(v):
    return round(Emu(v).inches, 2) if v is not None else None


def inspect(path):
    prs = Presentation(path)
    out = {"template": path,
           "slide_size_in": [emu_in(prs.slide_width), emu_in(prs.slide_height)],
           "layouts": []}
    for li, layout in enumerate(prs.slide_layouts):
        entry = {"index": li, "name": layout.name, "placeholders": []}
        for ph in layout.placeholders:
            f = ph.placeholder_format
            entry["placeholders"].append({
                "idx": f.idx,
                "type": str(f.type).split(" ")[0] if f.type is not None else "BODY",
                "name": ph.name,
                "pos_in": [emu_in(ph.left), emu_in(ph.top)],
                "size_in": [emu_in(ph.width), emu_in(ph.height)],
            })
        out["layouts"].append(entry)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("template")
    ap.add_argument("--json", help="機械可読の地図をこのパスに保存")
    args = ap.parse_args()

    data = inspect(args.template)
    print(f"テンプレ: {data['template']}  スライドサイズ: {data['slide_size_in']} in")
    for lay in data["layouts"]:
        print(f"\n[{lay['index']}] {lay['name']}")
        for ph in lay["placeholders"]:
            print(f"    idx={ph['idx']:<3} {ph['type']:<14} {ph['name']}  "
                  f"pos={ph['pos_in']} size={ph['size_in']}")
    if args.json:
        with open(args.json, "w", encoding="utf-8") as fp:
            json.dump(data, fp, ensure_ascii=False, indent=2)
        print(f"\n地図を保存: {args.json}")


if __name__ == "__main__":
    sys.exit(main())
